#!/usr/bin/env python3
"""
Convert the syndicate spreadsheet into JSON files for the web app.

Usage:
    python3 scripts/import-xlsx.py path/to/Copy_of_TWICHOL.xlsx

Writes:
    data/index.json          sheets, row/table counts, import timestamp, omitted columns
    data/<sheet-slug>.json   one file per sheet: { name, kind, tables: [{title, range, columns, rows}] }

Re-run whenever the spreadsheet changes, then commit the regenerated data/ files.

How sheets are read (the workbook is not a flat table, so layout is detected per sheet):
  * "tracker"  any sheet with a "Week total" row (the weekly pot tracker). Its blocks are found by their
               label cells (Punts/Bets, Fines Paid, Banked, Leaderboard Source), not fixed row numbers,
               so header rows that are not row 1 and blocks that move are handled. Every other labelled
               figure on the sheet is written to a "Summary cells" table with its cell address and formula.
  * "tournament" the 6-a-side organiser sheet (a club list plus a task list with a forward-filled task column).
  * "grid"     anything else: the sheet is split into blocks at blank rows and at gaps of two or more blank
               columns, the first text row of each block is used as its header, and blank columns are dropped.
  * empty sheets are kept in the index with zero tables.

Values: formula results are read from the cached values (what Excel last calculated). Dates and times are
written as ISO strings, floats are rounded to 6 dp, leading/trailing spaces and non-breaking spaces are removed.

Privacy: columns whose headers look like personal, contact or banking data (bank/sort code/account numbers,
phone, e-mail, contact names, addresses, dates of birth, passport, NI number, passwords, card details) are
dropped and listed in data/index.json under "omittedColumns". Phone numbers and e-mail addresses that appear
inside free text are replaced with "[phone removed]" / "[email removed]". Edit SENSITIVE_HEADER to change this.

Note: password-protected or rights-managed .xlsx files (e.g. Microsoft IRM / AIP-protected) cannot be read.
Save an unprotected copy first.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sys
import zipfile
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"

SENSITIVE_HEADER = re.compile(
    r"bank|sort\s*code|account\s*(no|num|number)|iban|\bbic\b|phone|mobile|tel\b|"
    r"e-?mail|address|postcode|post\s*code|\bdob\b|birth|passport|national\s*insurance|"
    r"password|\bpin\b|card\s*(no|number)|\bcvv\b|contact",
    re.I,
)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
# UK style numbers: 07..., (0)7..., (07...), +44 ... Only redacted when they contain 10+ digits.
PHONE_RE = re.compile(r"(?:\+44\s?\(?0?\)?\s?|\(0\)\s?|(?<![\w.])0)\d[\d\s().-]{8,}\d")
WEEK_RE = re.compile(r"^\s*week\s*(\d+)\s*$", re.I)


def is_plain_xlsx(path: Path) -> bool:
    """True only for a normal (unencrypted) OOXML zip package."""
    if not zipfile.is_zipfile(path):
        return False
    with zipfile.ZipFile(path) as z:
        return "xl/workbook.xml" in z.namelist()


def slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return slug or "sheet"


def scrub(text: str) -> str:
    text = EMAIL_RE.sub("[email removed]", text)

    def redact(m: re.Match) -> str:
        digits = sum(ch.isdigit() for ch in m.group(0))
        return "[phone removed]" if digits >= 10 else m.group(0)

    return PHONE_RE.sub(redact, text)


def norm(v):
    """Strip spaces and non-breaking spaces; empty strings become None."""
    if isinstance(v, str):
        s = v.replace("\xa0", " ").strip()
        return s or None
    return v


def to_json(v):
    v = norm(v)
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, int):
        return v
    if isinstance(v, str):
        return scrub(v)
    return scrub(str(v))


class Sheet:
    """Thin wrapper: normalised cell access on the cached-values sheet plus the formula sheet."""

    def __init__(self, ws_values, ws_formulas):
        self.v = ws_values
        self.f = ws_formulas
        self.max_row = ws_values.max_row
        self.max_column = ws_values.max_column

    def val(self, r: int, c: int):
        if r < 1 or c < 1 or r > self.max_row or c > self.max_column:
            return None
        return norm(self.v.cell(row=r, column=c).value)

    def text(self, r: int, c: int):
        v = self.val(r, c)
        return v.strip() if isinstance(v, str) else None

    def formula(self, r: int, c: int):
        f = self.f.cell(row=r, column=c).value
        return f[1:] if isinstance(f, str) and f.startswith("=") else None

    def find_text(self, pred):
        for r in range(1, self.max_row + 1):
            for c in range(1, self.max_column + 1):
                t = self.text(r, c)
                if t is not None and pred(t):
                    yield r, c

    def row_is_blank(self, r: int, c0: int = 1, c1: int | None = None) -> bool:
        c1 = c1 or self.max_column
        return all(self.val(r, c) is None for c in range(c0, c1 + 1))


def canon_week(header: str) -> str:
    m = WEEK_RE.match(header)
    return f"Week {int(m.group(1))}" if m else header


def make_table(title, range_ref, columns, rows):
    return {"title": title, "range": range_ref, "columns": columns, "rows": rows}


def drop_sensitive(columns: list[str], rows: list[dict], omitted: list[str]):
    """Remove columns whose header looks personal or banking; record their names in `omitted`."""
    keep = []
    for c in columns:
        if SENSITIVE_HEADER.search(c):
            omitted.append(c)
        else:
            keep.append(c)
    return keep, [{c: r.get(c) for c in keep} for r in rows]


# ---------------------------------------------------------------- tracker sheets

def week_block(sh: Sheet, hr: int, lc: int, title: str):
    """Block whose header row (hr) has the label in column lc and 'Week n' headers to the right."""
    weeks = []
    for c in range(lc + 1, sh.max_column + 1):
        t = sh.text(hr, c)
        if t and WEEK_RE.match(t):
            weeks.append((c, canon_week(t)))
    columns = ["Member"] + [name for _, name in weeks]
    rows, r = [], hr + 1
    while r <= sh.max_row:
        label = sh.text(r, lc)
        if label is None:
            break
        row = {"Member": to_json(label)}
        for c, name in weeks:
            row[name] = to_json(sh.val(r, c))
        rows.append(row)
        if label.lower().startswith("week total"):
            break
        r += 1
    rng = f"{get_column_letter(lc)}{hr}:{get_column_letter(weeks[-1][0])}{r}" if weeks else f"row {hr}"
    return make_table(title, rng, columns, rows), (hr, r, lc, weeks[-1][0] if weeks else lc)


def labelled_block(sh: Sheet, hr: int, lc: int, title: str, stop_prefix: str):
    """Block with a label column (lc) and a run of text headers starting at lc + 1."""
    heads = []
    c = lc + 1
    while sh.text(hr, c):
        heads.append(sh.text(hr, c))
        c += 1
    columns = ["Member"] + heads
    rows, r = [], hr + 1
    while True:
        label = sh.text(r, lc)
        if label is None:
            break
        row = {"Member": to_json(label)}
        for i, h in enumerate(heads):
            row[h] = to_json(sh.val(r, lc + 1 + i))
        rows.append(row)
        if label.lower().startswith(stop_prefix):
            break
        r += 1
    rng = f"{get_column_letter(lc)}{hr}:{get_column_letter(lc + len(heads))}{r}"
    return make_table(title, rng, columns, rows), (hr, r, lc, lc + len(heads))


def read_tracker(sh: Sheet):
    tables, claimed = [], []
    omitted: list[str] = []

    # Weekly punts: header label "Punts" (new layout) or "Bets" (old layout), followed by "Week 1".
    for r, c in sh.find_text(lambda s: s.lower() in ("punts", "bets")):
        if WEEK_RE.match(sh.text(r, c + 1) or ""):
            t, box = week_block(sh, r, c, "Weekly punts")
            tables.append(t)
            claimed.append(box)
            break

    # Banked: label "Banked" followed by "Week 1".
    for r, c in sh.find_text(lambda s: s.lower() == "banked"):
        if WEEK_RE.match(sh.text(r, c + 1) or ""):
            t, box = week_block(sh, r, c, "Banked")
            tables.append(t)
            claimed.append(box)
            break

    # Fines: header "Fines Paid" with "Overdue" beside it; label column is to its left.
    for r, c in sh.find_text(lambda s: s.lower().startswith("fines paid")):
        t, box = labelled_block(sh, r, c - 1, "Fines", "total")
        tables.append(t)
        claimed.append(box)
        break

    # Leaderboard source: header "Leaderboard Source" with "Total" and "Level".
    for r, c in sh.find_text(lambda s: s.lower() == "leaderboard source"):
        t, box = labelled_block(sh, r, c, "Leaderboard source", "\0")
        tables.append(t)
        claimed.append(box)
        break

    for t in tables:
        cols, rows = drop_sensitive(t["columns"], t["rows"], omitted)
        t["columns"], t["rows"] = cols, rows

    # Every other labelled figure (totals, interest, dates, targets) with its address and formula.
    def in_claimed(r, c):
        return any(r0 <= r <= r1 and c0 <= c <= c1 for r0, r1, c0, c1 in claimed)

    summary_rows = []
    for r in range(1, sh.max_row + 1):
        for c in range(1, sh.max_column + 1):
            label = sh.text(r, c)
            if label is None or in_claimed(r, c):
                continue
            for vc in range(c + 1, min(c + 4, sh.max_column) + 1):
                if in_claimed(r, vc):
                    continue
                value = sh.val(r, vc)
                if value is None:
                    continue
                summary_rows.append({
                    "Cell": f"{get_column_letter(c)}{r}",
                    "Label": to_json(label),
                    "Value": to_json(value),
                    "Formula": sh.formula(r, vc),
                })
                break
    if summary_rows:
        tables.append(make_table("Summary cells", "free cells", ["Cell", "Label", "Value", "Formula"], summary_rows))
    return tables, omitted


# ---------------------------------------------------------------- tournament sheet

def read_tournament(sh: Sheet):
    """Club list (header row with 'Club') and a task list (header row with 'TASK')."""
    tables, omitted = [], []
    for r, c in sh.find_text(lambda s: s.lower() == "club"):
        heads, cc = [], 1
        while sh.text(r, cc):
            heads.append(sh.text(r, cc))
            cc += 1
        rows, rr = [], r + 1
        while not sh.row_is_blank(rr, 1, len(heads)):
            rows.append({h: to_json(sh.val(rr, i + 1)) for i, h in enumerate(heads)})
            rr += 1
        cols, rows = drop_sensitive(heads, rows, omitted)
        tables.append(make_table("Clubs and teams", f"A{r}:{get_column_letter(len(heads))}{rr - 1}", cols, rows))
        break

    for r, c in sh.find_text(lambda s: s.lower() == "task"):
        title = next((sh.text(k, c) for k in range(r - 1, 0, -1) if sh.text(k, c)), "Tasks")
        heads = [sh.text(r, c + i) or f"Column {i + 1}" for i in range(3)]
        cols = ["Task", "Item / query / notes", "Responsibility"]
        rows, task, rr = [], None, r + 1
        while rr <= sh.max_row:
            task = sh.text(rr, c) or task
            item, resp = sh.val(rr, c + 1), sh.val(rr, c + 2)
            if item is not None or resp is not None:
                rows.append({"Task": to_json(task), cols[1]: to_json(item), cols[2]: to_json(resp)})
            rr += 1
        tables.append(make_table(title, f"{get_column_letter(c)}{r}:{get_column_letter(c + 2)}{rr - 1}", cols, rows))
        break
    return tables, omitted


# ---------------------------------------------------------------- generic grid sheets

def read_grid(sh: Sheet):
    """Split into blocks at blank rows and at gaps of 2+ blank columns; first text row is the header."""
    omitted: list[str] = []
    tables = []
    occupied_rows = [r for r in range(1, sh.max_row + 1) if not sh.row_is_blank(r)]
    bands, band = [], []
    for r in occupied_rows:
        if band and r != band[-1] + 1:
            bands.append(band)
            band = []
        band.append(r)
    if band:
        bands.append(band)

    for band in bands:
        occupied_cols = [c for c in range(1, sh.max_column + 1)
                         if any(sh.val(r, c) is not None for r in band)]
        runs, run = [], []
        for c in occupied_cols:
            if run and c - run[-1] > 2:
                runs.append(run)
                run = []
            run.append(c)
        if run:
            runs.append(run)

        for run in runs:
            c0, c1 = run[0], run[-1]
            body = [r for r in band if any(sh.val(r, c) is not None for c in range(c0, c1 + 1))]
            if not body:
                continue
            # Only the first row of the band can be a header, and only if it has 2+ values in this block.
            first = band[0]
            head_vals = [sh.val(first, c) for c in range(c0, c1 + 1)]
            has_header = sum(v is not None for v in head_vals) >= 2
            data_rows = [r for r in body if r != first] if has_header else body

            names = []
            for c in range(c0, c1 + 1):
                h = sh.val(first, c) if has_header else None
                if isinstance(h, float) and h.is_integer():
                    h = int(h)
                names.append(str(h).strip() if h is not None else f"Column {get_column_letter(c)}")
            keep = [i for i, c in enumerate(range(c0, c1 + 1))
                    if any(sh.val(r, c) is not None for r in data_rows)
                    or (has_header and sh.val(first, c) is not None)]
            cols, seen = [], {}
            for i in keep:
                base = names[i]
                n = seen.get(base, 0) + 1
                seen[base] = n
                cols.append(base if n == 1 else f"{base} ({n})")
            rows = [{cols[j]: to_json(sh.val(r, c0 + i)) for j, i in enumerate(keep)} for r in data_rows]
            cols, rows = drop_sensitive(cols, rows, omitted)
            rng = f"{get_column_letter(c0)}{band[0]}:{get_column_letter(c1)}{band[-1]}"
            tables.append(make_table(f"Table {rng}", rng, cols, rows))
    return tables, omitted


def read_sheet(sh: Sheet):
    if any(True for _ in sh.find_text(lambda s: s.lower() == "week total")):
        tables, omitted = read_tracker(sh)
        return "tracker", tables, omitted
    if any(True for _ in sh.find_text(lambda s: s.lower() == "task")):
        tables, omitted = read_tournament(sh)
        return "tournament", tables, omitted
    tables, omitted = read_grid(sh)
    return ("grid" if tables else "empty"), tables, omitted


def count_uncached_formulas(ws_values, ws_formulas) -> int:
    n = 0
    for row in ws_formulas.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                if ws_values[cell.coordinate].value is None:
                    n += 1
    return n


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2

    src = Path(sys.argv[1]).expanduser()
    if not src.exists():
        print(f"File not found: {src}", file=sys.stderr)
        return 1

    if not is_plain_xlsx(src):
        print(
            f"{src} is not a readable .xlsx file. It is probably password-protected or\n"
            "rights-managed (Microsoft IRM/AIP). Open it in Excel, remove the protection\n"
            "(File > Info > Protect Workbook / Restrict Access) and Save As a new .xlsx.",
            file=sys.stderr,
        )
        return 1

    wb_values = openpyxl.load_workbook(src, data_only=True)
    wb_formulas = openpyxl.load_workbook(src, data_only=False)
    DATA_DIR.mkdir(exist_ok=True)

    index = {
        "source": src.name,
        "importedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "sheets": [],
        "omittedColumns": {},
        "formulasWithoutCachedValue": 0,
    }

    written: set[str] = set()
    used_slugs: set[str] = set()
    for ws_v in wb_values.worksheets:
        ws_f = wb_formulas[ws_v.title]
        index["formulasWithoutCachedValue"] += count_uncached_formulas(ws_v, ws_f)
        sh = Sheet(ws_v, ws_f)
        kind, tables, omitted = read_sheet(sh)

        slug = slugify(ws_v.title)
        base, n = slug, 2
        while slug in used_slugs:
            slug = f"{base}-{n}"
            n += 1
        used_slugs.add(slug)

        payload = {"name": ws_v.title.strip(), "kind": kind, "tables": tables}
        (DATA_DIR / f"{slug}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        written.add(f"{slug}.json")
        row_count = sum(len(t["rows"]) for t in tables)
        col_count = max((len(t["columns"]) for t in tables), default=0)
        index["sheets"].append({
            "name": ws_v.title.strip(),
            "slug": slug,
            "kind": kind,
            "tableCount": len(tables),
            "rowCount": row_count,
            "columnCount": col_count,
        })
        if omitted:
            index["omittedColumns"][ws_v.title.strip()] = omitted
        print(f"{ws_v.title!r}: {kind}, {len(tables)} tables, {row_count} rows -> data/{slug}.json")

    # Remove JSON for sheets that no longer exist in the workbook.
    for stale in DATA_DIR.glob("*.json"):
        if stale.name != "index.json" and stale.name not in written:
            stale.unlink()
            print(f"Removed stale {stale.name}")

    (DATA_DIR / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    if index["formulasWithoutCachedValue"]:
        print(f"Warning: {index['formulasWithoutCachedValue']} formula cells have no cached value. "
              "Open and save the workbook in Excel, then re-run.")
    print("Wrote data/index.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
